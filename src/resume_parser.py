"""Extract plain text from an uploaded resume file (PDF, DOCX, or TXT)."""

import os

from pypdf import PdfReader
from docx import Document


def parse_resume(file_path: str) -> str:
    """Return the plain-text contents of a resume file.

    Supports .pdf, .docx, and .txt. Raises ValueError for anything else.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _parse_pdf(file_path)
    if ext == ".docx":
        return _parse_docx(file_path)
    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    raise ValueError(f"Unsupported resume format: {ext}. Please upload a PDF, DOCX, or TXT file.")


def _parse_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if not text:
        raise ValueError(
            "Could not extract any text from this PDF. It may be a scanned image "
            "without a text layer — try exporting a text-based PDF instead."
        )
    return text


def _parse_docx(file_path: str) -> str:
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    text = "\n".join(parts).strip()
    if not text:
        raise ValueError("Could not extract any text from this DOCX file.")
    return text
